"""
test_rpa_workflow.py
====================
Integration and unit tests for src/rpa_workflow.py and related modules.

Uses temporary directories and generated test workbooks — no user files required.
"""

from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

import pandas as pd
import pytest

from src.excel_reader import _normalise_col, read_workbook
from src.file_discovery import discover_files
from src.exception_handler import EmptyDirectoryError, CorruptedFileError
from src.performance_analyzer import PerformanceMetrics
from src.rpa_workflow import RPAWorkflow, WorkflowStatus


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _write_xlsx(path: Path, df: pd.DataFrame) -> None:
    df.to_excel(path, index=False, engine="openpyxl")


def _make_config(input_dir: str, output_dir: str) -> dict:
    return {
        "input_directory":         input_dir,
        "output_directory":        output_dir,
        "archive_directory":       output_dir + "/archive",
        "logs_directory":          output_dir + "/logs",
        "reports_directory":       output_dir + "/reports",
        "output_filename_prefix":  "Test_Master",
        "column_aliases": {
            "record_id": ["id", "emp_id", "record_id"],
            "name":      ["full_name", "employee_name", "name"],
            "email":     ["email_address", "email"],
            "salary":    ["annual_salary", "salary"],
        },
        "required_columns":         ["record_id", "name", "email", "salary"],
        "unique_key_columns":       ["record_id"],
        "duplicate_check_columns":  ["record_id"],
        "categorical_mappings":     {},
        "missing_value_policy": {
            "record_id": "flag",
            "name":      "flag",
            "email":     "flag",
            "salary":    "flag",
        },
        "validation_rules": {
            "salary_min":         0,
            "salary_max":         10000000,
            "email_regex":        "^[\\w\\.\\-]+@[\\w\\.\\-]+\\.\\w{2,}$",
            "date_formats":       ["%Y-%m-%d", "%d/%m/%Y"],
            "output_date_format": "%Y-%m-%d",
        },
        "performance": {"manual_records_per_minute": 100},
    }


@pytest.fixture
def tmp_dirs():
    root = tempfile.mkdtemp()
    inp  = Path(root) / "input"
    out  = Path(root) / "output"
    inp.mkdir()
    out.mkdir()
    yield inp, out
    shutil.rmtree(root)


@pytest.fixture
def clean_xlsx(tmp_dirs):
    inp, out = tmp_dirs
    df = pd.DataFrame({
        "record_id":    ["EMP001", "EMP002", "EMP003"],
        "name":         ["Alice Smith", "Bob Jones", "Carol White"],
        "email":        ["alice@example.com", "bob@example.com", "carol@example.com"],
        "salary":       [50000.0, 60000.0, 70000.0],
        "department":   ["Finance", "HR", "IT"],
        "joining_date": ["2020-01-15", "2019-03-15", "2021-07-01"],
        "city":         ["New York", "Chicago", "Boston"],
    })
    path = inp / "clean.xlsx"
    _write_xlsx(path, df)
    return path, inp, out


# ---------------------------------------------------------------------------
# 1. Column-name normalisation
# ---------------------------------------------------------------------------

class TestColumnNormalisation:
    def test_spaces_to_underscore(self):
        assert _normalise_col("Record ID") == "record_id"

    def test_already_clean(self):
        assert _normalise_col("salary") == "salary"

    def test_mixed_case_lowered(self):
        assert _normalise_col("Employee_Name") == "employee_name"

    def test_dots_to_underscore(self):
        assert _normalise_col("joining.date") == "joining_date"

    def test_multiple_spaces_collapsed(self):
        assert _normalise_col("  Full  Name  ") == "full_name"


# ---------------------------------------------------------------------------
# 11. Empty input directory
# ---------------------------------------------------------------------------

class TestEmptyDirectory:
    def test_empty_dir_raises(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            with pytest.raises(EmptyDirectoryError):
                discover_files(tmpdir)

    def test_missing_dir_raises(self):
        from src.exception_handler import FileDiscoveryError
        with pytest.raises(FileDiscoveryError):
            discover_files("/nonexistent/path/xyz")


# ---------------------------------------------------------------------------
# 12. Corrupted file handling
# ---------------------------------------------------------------------------

class TestCorruptedFile:
    def test_corrupted_file_raises(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            bad = Path(tmpdir) / "corrupt.xlsx"
            bad.write_bytes(b"this is not an xlsx file at all!!!")
            with pytest.raises(CorruptedFileError):
                read_workbook(bad)


# ---------------------------------------------------------------------------
# 13. Missing mandatory columns
# ---------------------------------------------------------------------------

class TestMissingMandatoryColumns:
    def test_missing_required_col_in_validator(self):
        from src.data_validator import validate_dataframe
        df = pd.DataFrame({
            "record_id": ["EMP001"],
            # 'name' intentionally absent
            "email":     ["alice@test.com"],
            "salary":    [50000.0],
        })
        report = validate_dataframe(df, required_columns=["name"])
        # Missing column → all records flagged critical
        critical = [e for e in report.errors if e.severity == "critical"]
        assert len(critical) >= 1


# ---------------------------------------------------------------------------
# 14. Output write failure handling
# ---------------------------------------------------------------------------

class TestOutputWriteFailure:
    def test_write_to_nonexistent_parent_raises(self):
        from src.excel_exporter import export_master_workbook
        from src.exception_handler import OutputWriteError
        import os, stat, tempfile
        with tempfile.TemporaryDirectory() as tmpdir:
            # Create a read-only subdirectory so writing inside it fails
            locked = os.path.join(tmpdir, "locked")
            os.makedirs(locked)
            os.chmod(locked, stat.S_IRUSR | stat.S_IXUSR)  # read+exec only, no write
            bad_path = os.path.join(locked, "subdir", "out.xlsx")
            try:
                with pytest.raises(OutputWriteError):
                    export_master_workbook(
                        output_path    = bad_path,
                        master_df      = pd.DataFrame({"a": [1]}),
                        exceptions_df  = pd.DataFrame(),
                        duplicates_df  = pd.DataFrame(),
                        processing_log = [],
                        cleaning_summary=[],
                        run_stats      = {},
                    )
            finally:
                # Restore so cleanup can delete it
                os.chmod(locked, stat.S_IRWXU)


# ---------------------------------------------------------------------------
# 15. Original input files preserved
# ---------------------------------------------------------------------------

class TestInputFilePreservation:
    def test_original_file_unchanged(self, clean_xlsx):
        path, inp, out = clean_xlsx
        import hashlib
        original_hash = hashlib.md5(path.read_bytes()).hexdigest()

        config = _make_config(str(inp), str(out))
        wf = RPAWorkflow(config)
        wf.run()

        new_hash = hashlib.md5(path.read_bytes()).hexdigest()
        assert original_hash == new_hash, "Input file was modified!"


# ---------------------------------------------------------------------------
# 16. Performance calculations
# ---------------------------------------------------------------------------

class TestPerformanceCalculations:
    def test_basic_calculation(self):
        m = PerformanceMetrics(
            total_input_records    = 600,
            automation_duration_s  = 10.0,
            manual_rate_per_minute = 100,
        ).calculate()
        # manual time = 600 / (100/60) = 360 s
        assert abs(m.estimated_manual_duration_s - 360.0) < 0.1
        assert m.estimated_time_saved_s == pytest.approx(350.0, abs=0.5)
        assert m.records_per_second == pytest.approx(60.0, abs=0.1)

    def test_zero_records_safe(self):
        m = PerformanceMetrics(
            total_input_records   = 0,
            automation_duration_s = 5.0,
        ).calculate()
        assert m.records_per_second == 0.0
        assert m.estimated_manual_duration_s == 0.0

    def test_zero_duration_safe(self):
        m = PerformanceMetrics(
            total_input_records   = 100,
            automation_duration_s = 0.0,
        ).calculate()
        # Should not raise ZeroDivisionError
        assert m.records_per_second == float("inf") or m.records_per_second >= 0

    def test_negative_time_saved_reported_accurately(self):
        m = PerformanceMetrics(
            total_input_records    = 1,   # tiny dataset
            automation_duration_s  = 100.0,
            manual_rate_per_minute = 100,
        ).calculate()
        # Manual time = 1 / (100/60) = 0.6 s → saved is negative
        assert m.estimated_time_saved_s < 0


# ---------------------------------------------------------------------------
# 17. End-to-end workflow
# ---------------------------------------------------------------------------

class TestEndToEndWorkflow:
    def test_workflow_succeeds_on_clean_file(self, clean_xlsx):
        _, inp, out = clean_xlsx
        config = _make_config(str(inp), str(out))
        wf = RPAWorkflow(config)
        result = wf.run()

        assert result.status in (WorkflowStatus.SUCCESS, WorkflowStatus.PARTIAL_SUCCESS)
        assert result.output_workbook is not None
        assert Path(result.output_workbook).exists()

    def test_master_data_not_empty(self, clean_xlsx):
        _, inp, out = clean_xlsx
        config = _make_config(str(inp), str(out))
        wf = RPAWorkflow(config)
        result = wf.run()
        assert not result.master_df.empty

    def test_expected_worksheets_in_workbook(self, clean_xlsx):
        _, inp, out = clean_xlsx
        config = _make_config(str(inp), str(out))
        wf = RPAWorkflow(config)
        result = wf.run()

        from openpyxl import load_workbook
        wb = load_workbook(result.output_workbook, read_only=True)
        sheets = wb.sheetnames
        wb.close()
        expected = {"Master_Data", "Data_Quality_Report", "Exceptions",
                    "Processing_Log", "Duplicate_Records", "Run_Summary"}
        for sheet in expected:
            assert sheet in sheets, f"Missing worksheet: {sheet}"

    def test_failed_file_reported(self, tmp_dirs):
        inp, out = tmp_dirs
        bad = inp / "corrupt.xlsx"
        bad.write_bytes(b"NOT A VALID XLSX FILE")
        config = _make_config(str(inp), str(out))
        wf = RPAWorkflow(config)
        result = wf.run()
        # Should be FAILED (only file was corrupt) or PARTIAL_SUCCESS
        assert result.status in (WorkflowStatus.FAILED, WorkflowStatus.PARTIAL_SUCCESS)
        failed = [r for r in result.file_results if r["status"] == "FAILED"]
        assert len(failed) >= 1

    def test_metrics_populated(self, clean_xlsx):
        _, inp, out = clean_xlsx
        config = _make_config(str(inp), str(out))
        wf = RPAWorkflow(config)
        result = wf.run()
        assert result.metrics is not None
        assert result.metrics.total_input_records > 0
