"""
excel_exporter.py
=================
Writes the master consolidated workbook with professional Excel formatting.

Worksheets generated
---------------------
1. Master_Data           – Consolidated, cleaned records.
2. Data_Quality_Report   – Cleaning statistics.
3. Exceptions            – Records requiring review / invalid records.
4. Processing_Log        – Per-file processing details.
5. Duplicate_Records     – Detected exact and conflicting duplicates.
6. Run_Summary           – Overall execution statistics.

Design principles
-----------------
* Atomic write: write to a temp file, then rename — never corrupt a good output.
* Never overwrite existing output silently; caller passes a unique filename.
* Professional formatting: frozen headers, auto-width, colour coding.
"""

from __future__ import annotations

import logging
import os
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from src.exception_handler import OutputWriteError

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------

HEADER_FILL   = PatternFill("solid", fgColor="1F3864")   # dark navy
HEADER_FONT   = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
ALT_ROW_FILL  = PatternFill("solid", fgColor="E9EFF7")   # light blue
VALID_FILL    = PatternFill("solid", fgColor="C6EFCE")   # light green
REVIEW_FILL   = PatternFill("solid", fgColor="FFEB9C")   # light yellow
INVALID_FILL  = PatternFill("solid", fgColor="FFC7CE")   # light red
SUCCESS_FILL  = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL     = PatternFill("solid", fgColor="FFC7CE")
THIN_BORDER   = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def export_master_workbook(
    output_path: str | Path,
    master_df: pd.DataFrame,
    exceptions_df: pd.DataFrame,
    duplicates_df: pd.DataFrame,
    processing_log: list[dict],
    cleaning_summary: list[dict],
    run_stats: dict,
) -> Path:
    """
    Write all worksheets to *output_path* using an atomic write pattern.

    Returns the final output path.

    Raises
    ------
    OutputWriteError – on any write failure (original file preserved).
    """
    output_path = Path(output_path)

    # Write to a temp file first, then rename atomically
    tmp_path: Path | None = None
    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)

        fd, tmp = tempfile.mkstemp(
            suffix=".xlsx", dir=output_path.parent, prefix="_tmp_"
        )
        os.close(fd)
        tmp_path = Path(tmp)

        wb = Workbook()
        wb.remove(wb.active)  # remove default empty sheet

        _write_master_data(wb, master_df)
        _write_data_quality(wb, cleaning_summary)
        _write_exceptions(wb, exceptions_df)
        _write_processing_log(wb, processing_log)
        _write_duplicates(wb, duplicates_df)
        _write_run_summary(wb, run_stats)

        wb.save(str(tmp_path))
        tmp_path.replace(output_path)   # atomic rename
        logger.info("Master workbook written: %s", output_path)
        return output_path

    except Exception as exc:
        # Clean up temp file; leave original output untouched
        if tmp_path and tmp_path.exists():
            try:
                tmp_path.unlink()
            except OSError:
                pass
        raise OutputWriteError(
            f"Failed to write master workbook to '{output_path}': {exc}",
            stage="Excel Export",
        ) from exc


# ---------------------------------------------------------------------------
# Sheet writers
# ---------------------------------------------------------------------------

def _write_master_data(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Master_Data")
    if df.empty:
        ws["A1"] = "No records to display."
        return
    _df_to_sheet(ws, df, status_col="_validation_status")


def _write_data_quality(wb: Workbook, cleaning_summary: list[dict]) -> None:
    ws = wb.create_sheet("Data_Quality_Report")
    if not cleaning_summary:
        ws["A1"] = "No cleaning operations were recorded."
        return

    rows: list[dict] = []
    for entry in cleaning_summary:
        source = entry.get("source_file", "")
        for op in entry.get("operations", []):
            rows.append({
                "Source File":         source,
                "Operation":           op.get("operation", ""),
                "Column":              op.get("column", ""),
                "Records Affected":    op.get("affected", 0),
                "Detail":              op.get("detail", ""),
            })

    df = pd.DataFrame(rows) if rows else pd.DataFrame(
        columns=["Source File", "Operation", "Column", "Records Affected", "Detail"]
    )
    _df_to_sheet(ws, df)


def _write_exceptions(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Exceptions")
    if df.empty:
        ws["A1"] = "No exceptions recorded."
        return
    _df_to_sheet(ws, df, status_col="_validation_status")


def _write_processing_log(wb: Workbook, log: list[dict]) -> None:
    ws = wb.create_sheet("Processing_Log")
    if not log:
        ws["A1"] = "No processing log entries."
        return
    df = pd.DataFrame(log)
    _df_to_sheet(ws, df, status_col="status")


def _write_duplicates(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Duplicate_Records")
    if df.empty:
        ws["A1"] = "No duplicate records detected."
        return
    _df_to_sheet(ws, df)


def _write_run_summary(wb: Workbook, stats: dict) -> None:
    ws = wb.create_sheet("Run_Summary")

    headers = ["Metric", "Value"]
    rows = list(stats.items())

    # Header row
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill  = HEADER_FILL
        cell.font  = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    for row_idx, (metric, value) in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=str(metric)).border = THIN_BORDER
        ws.cell(row=row_idx, column=2, value=str(value)).border = THIN_BORDER

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 30
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Generic DataFrame → worksheet helper
# ---------------------------------------------------------------------------

def _df_to_sheet(
    ws,
    df: pd.DataFrame,
    status_col: str | None = None,
) -> None:
    """Write a DataFrame to *ws* with formatting."""
    # Remove internal meta columns from display
    display_df = df.copy()
    hide_cols = [c for c in display_df.columns if c.startswith("_missing_")]
    display_df = display_df.drop(columns=hide_cols, errors="ignore")

    # Rename underscore columns to readable headers
    display_df.columns = [
        _col_to_header(c) for c in display_df.columns
    ]

    rows = list(dataframe_to_rows(display_df, index=False, header=True))
    if not rows:
        ws["A1"] = "Empty."
        return

    header_row = rows[0]
    status_col_display = _col_to_header(status_col) if status_col else None

    for col_idx, header in enumerate(header_row, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill      = HEADER_FILL
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = THIN_BORDER

    for row_idx, row_data in enumerate(rows[1:], start=2):
        fill = ALT_ROW_FILL if row_idx % 2 == 0 else None

        # Determine status-based fill
        if status_col_display and status_col_display in header_row:
            sc_idx = header_row.index(status_col_display)
            status_val = str(row_data[sc_idx]).upper() if sc_idx < len(row_data) else ""
            if "VALID" in status_val and "INVALID" not in status_val:
                fill = VALID_FILL
            elif "REVIEW" in status_val:
                fill = REVIEW_FILL
            elif "INVALID" in status_val or "FAIL" in status_val:
                fill = INVALID_FILL
            elif "SUCCESS" in status_val:
                fill = SUCCESS_FILL

        for col_idx, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if fill:
                cell.fill = fill

    # Auto-fit column widths (capped)
    for col_idx, header in enumerate(header_row, start=1):
        letter = get_column_letter(col_idx)
        max_len = max(
            len(str(header)),
            *[len(str(rows[i][col_idx - 1])) if col_idx - 1 < len(rows[i]) else 0
              for i in range(1, min(len(rows), 101))],
        )
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 50)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _col_to_header(col: str) -> str:
    """Convert snake_case column name to 'Title Case' header."""
    if col.startswith("_"):
        col = col.lstrip("_")
    return col.replace("_", " ").title()


def generate_output_filename(
    output_dir: str | Path,
    prefix: str = "Master_Consolidated",
) -> Path:
    """
    Return a unique output path:
    <output_dir>/<prefix>_YYYYMMDD_HHMMSS.xlsx
    """
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path(output_dir) / f"{prefix}_{ts}.xlsx"
